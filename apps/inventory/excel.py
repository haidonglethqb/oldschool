"""
Reusable Excel (.xlsx) upload parser, validator, preview generator, and safe exporter (Mục 12).
Supports:
- Import receipt lines (UC-4.1)
- Sales outbound lines (UC-5.2)
- Stocktake count lines (UC-6.3)
- Safe spreadsheet export with formula-injection protection (UC-4.10, UC-7.3)
"""
from decimal import Decimal, InvalidOperation
from io import BytesIO
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from django.conf import settings
from django.http import HttpResponse
from apps.branches.models import CoSo
from apps.catalog.models import BienThe
from apps.inventory.models import TonKho

FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def sanitize_excel_value(val):
    """Prevent spreadsheet formula injection when exporting untrusted text to Excel."""
    if isinstance(val, str) and val.startswith(FORMULA_PREFIXES):
        return f"'{val}"
    return val


def validate_uploaded_xlsx(uploaded_file) -> tuple[openpyxl.Workbook | None, str | None]:
    """Validate file extension, size limit, and openpyxl parseability."""
    if not uploaded_file:
        return None, "Vui lòng chọn một file Excel (.xlsx) để tải lên."
    filename = getattr(uploaded_file, "name", "") or ""
    if not filename.lower().endswith(".xlsx"):
        return None, "Định dạng file không hợp lệ. Hệ thống chỉ hỗ trợ file Excel định dạng .xlsx."

    max_mb = getattr(settings, "MAX_UPLOAD_SIZE_MB", 10)
    max_bytes = max_mb * 1024 * 1024
    if getattr(uploaded_file, "size", 0) > max_bytes:
        return None, f"Dung lượng file vượt quá giới hạn cho phép ({max_mb} MB)."

    try:
        wb = openpyxl.load_workbook(uploaded_file, data_only=True)
        return wb, None
    except Exception:
        return None, "Không thể đọc nội dung file Excel. Vui lòng kiểm tra file không bị hỏng hoặc đặt mật khẩu."


def _normalize_header(val) -> str:
    if val is None:
        return ""
    s = str(val).strip().lower()
    mapping = {
        "mã sku": "ma_sku",
        "ma_sku": "ma_sku",
        "sku": "ma_sku",
        "mã hàng": "ma_sku",
        "mã hàng hóa": "ma_sku",
        "mã sản phẩm": "ma_sku",
        "ma_san_pham": "ma_sku",
        "ma_hang": "ma_sku",
        "số lượng": "so_luong",
        "so_luong": "so_luong",
        "số lượng nhập": "so_luong",
        "số lượng xuất": "so_luong",
        "số lượng thực tế": "so_luong_thuc_te",
        "so_luong_thuc_te": "so_luong_thuc_te",
        "thực tế": "so_luong_thuc_te",
        "đơn giá": "don_gia",
        "don_gia": "don_gia",
        "giá nhập": "don_gia",
        "giá bán": "don_gia",
        "chiết khấu": "chiet_khau",
        "chiet_khau": "chiet_khau",
        "ghi chú": "ghi_chu",
        "ghi_chu": "ghi_chu",
        "xử lý": "xu_ly",
        "xu_ly": "xu_ly",
    }
    return mapping.get(s, s)


def _resolve_variant(code: str) -> BienThe | None:
    clean_code = code.strip().upper()
    variant = (
        BienThe.objects.select_related("san_pham")
        .filter(ma_sku__iexact=clean_code, san_pham__trang_thai="ACTIVE", trang_thai="ACTIVE")
        .first()
    )
    if variant:
        return variant
    # Fallback if product has exactly 1 active variant matching ma_san_pham
    candidates = list(
        BienThe.objects.select_related("san_pham").filter(
            san_pham__ma_san_pham__iexact=clean_code,
            san_pham__trang_thai="ACTIVE",
            trang_thai="ACTIVE",
        )
    )
    if len(candidates) == 1:
        return candidates[0]
    return None


def parse_import_excel(uploaded_file) -> dict:
    """
    Parse and validate an Excel file for Warehouse Import (UC-4.1).
    Required columns: `ma_sku` (or `Mã SKU`/`Mã hàng`), `so_luong`, `don_gia`.
    Optional: `chiet_khau`, `ghi_chu`.
    """
    wb, err = validate_uploaded_xlsx(uploaded_file)
    if err:
        return {"valid": False, "file_error": err, "rows": [], "errors": [err]}

    ws = wb.active
    raw_rows = list(ws.iter_rows(values_only=True))
    if not raw_rows:
        return {"valid": False, "file_error": "File Excel trống.", "rows": [], "errors": ["File Excel trống."]}

    headers = [_normalize_header(c) for c in raw_rows[0]]
    required = {"ma_sku", "so_luong", "don_gia"}
    missing = required - set(headers)
    if missing:
        msg = "File thiếu cột bắt buộc. Cần có: Mã SKU (hoặc Mã hàng), Số lượng, Đơn giá."
        return {"valid": False, "file_error": msg, "rows": [], "errors": [msg]}

    col_idx = {name: headers.index(name) for name in headers if name}
    parsed_rows = []
    errors = []
    seen_skus = set()

    for idx, row in enumerate(raw_rows[1:], start=2):
        if not row or all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        sku_raw = str(row[col_idx["ma_sku"]] or "").strip()
        qty_raw = row[col_idx["so_luong"]]
        price_raw = row[col_idx["don_gia"]]
        disc_raw = row[col_idx["chiet_khau"]] if "chiet_khau" in col_idx and col_idx["chiet_khau"] < len(row) else 0
        note_raw = str(row[col_idx["ghi_chu"]] or "").strip() if "ghi_chu" in col_idx and col_idx["ghi_chu"] < len(row) else ""

        row_errors = []
        variant = _resolve_variant(sku_raw) if sku_raw else None
        if not sku_raw:
            row_errors.append("Thiếu Mã SKU/Mã hàng.")
        elif not variant:
            row_errors.append(f"Mã SKU/hàng '{sku_raw}' không tồn tại trong danh mục đang kinh doanh.")
        elif variant.ma_sku in seen_skus:
            row_errors.append(f"Mã SKU '{variant.ma_sku}' bị trùng lặp trong file.")
        else:
            seen_skus.add(variant.ma_sku)

        qty = 0
        try:
            qty = int(float(str(qty_raw)))
            if qty <= 0:
                row_errors.append("Số lượng nhập phải lớn hơn 0.")
        except (ValueError, TypeError):
            row_errors.append("Số lượng nhập không hợp lệ.")

        unit_price = Decimal("0")
        try:
            unit_price = Decimal(str(price_raw))
            if unit_price < 0:
                row_errors.append("Đơn giá nhập không được âm.")
        except (InvalidOperation, ValueError, TypeError):
            row_errors.append("Đơn giá nhập không hợp lệ.")

        discount = Decimal("0")
        if disc_raw not in (None, ""):
            try:
                discount = Decimal(str(disc_raw))
                if discount < 0:
                    row_errors.append("Chiết khấu không được âm.")
                elif qty > 0 and unit_price >= 0 and discount > Decimal(qty) * unit_price:
                    row_errors.append("Chiết khấu vượt quá thành tiền dòng hàng.")
            except (InvalidOperation, ValueError, TypeError):
                row_errors.append("Chiết khấu không hợp lệ.")

        line_total = max(Decimal("0"), Decimal(qty) * unit_price - discount) if not row_errors else Decimal("0")
        if row_errors:
            errors.append(f"Dòng {idx}: " + " ".join(row_errors))

        parsed_rows.append(
            {
                "row_number": idx,
                "sku_input": sku_raw,
                "variant_id": variant.id if variant else None,
                "ma_sku": variant.ma_sku if variant else sku_raw,
                "ten_hang_hoa": variant.full_display_name if variant else "Chưa xác định",
                "don_vi_tinh": variant.san_pham.don_vi_tinh if variant else "",
                "so_luong": qty,
                "don_gia": str(unit_price),
                "chiet_khau": str(discount),
                "thanh_tien": str(line_total),
                "ghi_chu": note_raw,
                "errors": row_errors,
            }
        )

    if not parsed_rows:
        errors.append("File Excel không có dòng dữ liệu hàng hóa nào.")

    return {
        "valid": len(errors) == 0 and len(parsed_rows) > 0,
        "file_error": None,
        "rows": parsed_rows,
        "errors": errors,
    }


def parse_outbound_excel(uploaded_file, branch: CoSo) -> dict:
    """
    Parse and validate an Excel file for Sales Outbound (UC-5.2).
    Required columns: `ma_sku` (or `Mã SKU`/`Mã hàng`), `so_luong`.
    Optional: `don_gia`, `ghi_chu`.
    """
    wb, err = validate_uploaded_xlsx(uploaded_file)
    if err:
        return {"valid": False, "file_error": err, "rows": [], "errors": [err]}

    ws = wb.active
    raw_rows = list(ws.iter_rows(values_only=True))
    if not raw_rows:
        return {"valid": False, "file_error": "File Excel trống.", "rows": [], "errors": ["File Excel trống."]}

    headers = [_normalize_header(c) for c in raw_rows[0]]
    required = {"ma_sku", "so_luong"}
    missing = required - set(headers)
    if missing:
        msg = "File thiếu cột bắt buộc. Cần có: Mã SKU (hoặc Mã hàng), Số lượng."
        return {"valid": False, "file_error": msg, "rows": [], "errors": [msg]}

    col_idx = {name: headers.index(name) for name in headers if name}
    parsed_rows = []
    errors = []
    seen_skus = set()

    for idx, row in enumerate(raw_rows[1:], start=2):
        if not row or all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        sku_raw = str(row[col_idx["ma_sku"]] or "").strip()
        qty_raw = row[col_idx["so_luong"]]
        price_raw = row[col_idx["don_gia"]] if "don_gia" in col_idx and col_idx["don_gia"] < len(row) else None
        note_raw = str(row[col_idx["ghi_chu"]] or "").strip() if "ghi_chu" in col_idx and col_idx["ghi_chu"] < len(row) else ""

        row_errors = []
        variant = _resolve_variant(sku_raw) if sku_raw else None
        if not sku_raw:
            row_errors.append("Thiếu Mã SKU/Mã hàng.")
        elif not variant:
            row_errors.append(f"Mã SKU/hàng '{sku_raw}' chưa có trong danh mục.")
        elif variant.ma_sku in seen_skus:
            row_errors.append(f"Mã SKU '{variant.ma_sku}' bị trùng lặp trong file.")
        else:
            seen_skus.add(variant.ma_sku)

        qty = 0
        try:
            qty = int(float(str(qty_raw)))
            if qty <= 0:
                row_errors.append("Số lượng xuất phải lớn hơn 0.")
        except (ValueError, TypeError):
            row_errors.append("Số lượng xuất không hợp lệ.")

        avail_qty = 0
        if variant and branch:
            bal = TonKho.objects.filter(co_so=branch, bien_the=variant).first()
            avail_qty = bal.so_luong if bal else 0
            if qty > avail_qty:
                row_errors.append(f"Số lượng xuất ({qty}) vượt quá tồn kho tại {branch.ma_co_so} ({avail_qty}).")

        unit_price = variant.don_gia if variant else Decimal("0")
        if price_raw not in (None, ""):
            try:
                unit_price = Decimal(str(price_raw))
                if unit_price < 0:
                    row_errors.append("Đơn giá bán không được âm.")
            except (InvalidOperation, ValueError, TypeError):
                row_errors.append("Đơn giá bán không hợp lệ.")

        line_total = Decimal(qty) * unit_price if not row_errors else Decimal("0")
        if row_errors:
            errors.append(f"Dòng {idx}: " + " ".join(row_errors))

        parsed_rows.append(
            {
                "row_number": idx,
                "variant_id": variant.id if variant else None,
                "ma_sku": variant.ma_sku if variant else sku_raw,
                "ten_hang_hoa": variant.full_display_name if variant else "Chưa xác định",
                "don_vi_tinh": variant.san_pham.don_vi_tinh if variant else "",
                "ton_hien_tai": avail_qty,
                "so_luong": qty,
                "don_gia": str(unit_price),
                "thanh_tien": str(line_total),
                "ghi_chu": note_raw,
                "errors": row_errors,
            }
        )

    if not parsed_rows:
        errors.append("File Excel không có dòng dữ liệu nào.")

    return {
        "valid": len(errors) == 0 and len(parsed_rows) > 0,
        "file_error": None,
        "rows": parsed_rows,
        "errors": errors,
    }


def parse_stocktake_excel(uploaded_file, branch: CoSo) -> dict:
    """
    Parse and validate an Excel file for Stocktake / Inventory Count (UC-6.3).
    Required columns: `ma_sku` (or `Mã hàng`), `so_luong_thuc_te` (or `Số lượng thực tế` / `Số lượng`).
    Optional: `xu_ly`, `ghi_chu`.
    """
    wb, err = validate_uploaded_xlsx(uploaded_file)
    if err:
        return {"valid": False, "file_error": err, "rows": [], "errors": [err]}

    ws = wb.active
    raw_rows = list(ws.iter_rows(values_only=True))
    if not raw_rows:
        return {"valid": False, "file_error": "File Excel trống.", "rows": [], "errors": ["File Excel trống."]}

    headers = [_normalize_header(c) for c in raw_rows[0]]
    qty_col_name = "so_luong_thuc_te" if "so_luong_thuc_te" in headers else ("so_luong" if "so_luong" in headers else "")
    if "ma_sku" not in headers or not qty_col_name:
        msg = "File thiếu dữ liệu bắt buộc. Cần có cột: Mã hàng (hoặc Mã SKU) và Số lượng thực tế."
        return {"valid": False, "file_error": msg, "rows": [], "errors": [msg]}

    col_idx = {name: headers.index(name) for name in headers if name}
    parsed_rows = []
    errors = []
    seen_skus = set()

    for idx, row in enumerate(raw_rows[1:], start=2):
        if not row or all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        sku_raw = str(row[col_idx["ma_sku"]] or "").strip()
        actual_raw = row[col_idx[qty_col_name]]
        xu_ly_raw = str(row[col_idx["xu_ly"]] or "").strip() if "xu_ly" in col_idx and col_idx["xu_ly"] < len(row) else ""
        note_raw = str(row[col_idx["ghi_chu"]] or "").strip() if "ghi_chu" in col_idx and col_idx["ghi_chu"] < len(row) else ""

        row_errors = []
        variant = _resolve_variant(sku_raw) if sku_raw else None
        if not sku_raw:
            row_errors.append("Thiếu Mã hàng/SKU.")
        elif not variant:
            row_errors.append(f"Mã hàng/SKU '{sku_raw}' không tồn tại tại cơ sở kiểm kê.")
        elif variant.ma_sku in seen_skus:
            row_errors.append(f"Mã hàng/SKU '{variant.ma_sku}' bị trùng lặp trong file.")
        else:
            seen_skus.add(variant.ma_sku)

        actual_qty = 0
        try:
            if actual_raw in (None, ""):
                raise ValueError("Empty")
            actual_qty = int(float(str(actual_raw)))
            if actual_qty < 0:
                row_errors.append("Số lượng thực tế không được nhỏ hơn 0.")
        except (ValueError, TypeError):
            row_errors.append("Số lượng thực tế không hợp lệ (phải là số nguyên >= 0).")

        sys_qty = 0
        if variant and branch:
            bal = TonKho.objects.filter(co_so=branch, bien_the=variant).first()
            sys_qty = bal.so_luong if bal else 0

        diff = actual_qty - sys_qty
        status_label = "Khớp" if diff == 0 else ("Thiếu" if diff < 0 else "Thừa")
        if row_errors:
            errors.append(f"Dòng {idx}: " + " ".join(row_errors))

        parsed_rows.append(
            {
                "row_number": idx,
                "variant_id": variant.id if variant else None,
                "ma_sku": variant.ma_sku if variant else sku_raw,
                "ten_hang_hoa": variant.full_display_name if variant else "Chưa xác định",
                "don_vi_tinh": variant.san_pham.don_vi_tinh if variant else "Cái",
                "don_gia": str(variant.don_gia if variant else Decimal("0")),
                "so_luong_he_thong": sys_qty,
                "so_luong_thuc_te": actual_qty,
                "chenh_lech": diff,
                "trang_thai_label": status_label,
                "xu_ly": xu_ly_raw,
                "ghi_chu": note_raw,
                "errors": row_errors,
            }
        )

    if not parsed_rows:
        errors.append("File Excel không có dòng dữ liệu kiểm kê nào.")

    return {
        "valid": len(errors) == 0 and len(parsed_rows) > 0,
        "file_error": None,
        "rows": parsed_rows,
        "errors": errors,
    }


def build_excel_response(*, filename: str, sheet_title: str, title_banner: str, subtitle: str, headers: list[str], rows: list[list]) -> HttpResponse:
    """Create a styled `.xlsx` download response with formula-injection protection."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    header_fill = PatternFill(start_color="4A2511", end_color="4A2511", fill_type="solid")
    sub_fill = PatternFill(start_color="DEC2A8", end_color="DEC2A8", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="4A2511")
    thin_border = Border(
        left=Side(style="thin", color="C8B6A6"),
        right=Side(style="thin", color="C8B6A6"),
        top=Side(style="thin", color="C8B6A6"),
        bottom=Side(style="thin", color="C8B6A6"),
    )

    ws.cell(row=1, column=1, value=sanitize_excel_value(title_banner)).font = title_font
    ws.cell(row=2, column=1, value=sanitize_excel_value(subtitle)).font = Font(italic=True, size=10)

    start_row = 4
    for col_num, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_num, value=sanitize_excel_value(h))
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for r_idx, row_values in enumerate(rows, start=start_row + 1):
        for c_idx, val in enumerate(row_values, start=1):
            c = ws.cell(row=r_idx, column=c_idx, value=sanitize_excel_value(val))
            c.border = thin_border

    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 4, 12), 42)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    response = HttpResponse(
        stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
