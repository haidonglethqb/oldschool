"""Regressions for downloaded spreadsheets, stocktake forms and counter-request roles."""
from html.parser import HTMLParser
from io import BytesIO
import unicodedata

import openpyxl
import pytest
from django.core.exceptions import PermissionDenied
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse

from apps.inventory.excel import parse_import_excel, parse_outbound_excel, parse_stocktake_excel
from apps.inventory.models import StockMovement, TonKho
from apps.notifications.models import ThongBao
from apps.purchasing.models import ChiTietPhieuNhapKho, PhieuNhapKho
from apps.sales_outbound.models import ChiTietPhieuXuatHangBan, ChiTietYeuCauXuatHang, PhieuXuatHangBan, YeuCauXuatHang
from apps.sales_outbound.services import OutboundRequestLineSpec, create_outbound_request
from apps.stocktakes.models import ChiTietKiemKe, KiemKe


pytestmark = pytest.mark.django_db


def workbook_upload(headers, rows, header_row=1):
    wb = openpyxl.Workbook()
    ws = wb.active
    if header_row > 1:
        ws.cell(1, 1, "OLD SCHOOL")
        ws.cell(2, 1, "Hướng dẫn điền file")
    for column, value in enumerate(headers, 1):
        ws.cell(header_row, column, value)
    for row_number, row in enumerate(rows, header_row + 1):
        for column, value in enumerate(row, 1):
            ws.cell(row_number, column, value)
    stream = BytesIO()
    wb.save(stream)
    return SimpleUploadedFile("counts.xlsx", stream.getvalue())


def parse(kind, upload, branch):
    if kind == "import":
        return parse_import_excel(upload)
    return {"outbound": parse_outbound_excel, "stocktake": parse_stocktake_excel}[kind](upload, branch)


def business_state():
    models = (PhieuNhapKho, ChiTietPhieuNhapKho, PhieuXuatHangBan, ChiTietPhieuXuatHangBan,
              KiemKe, ChiTietKiemKe, YeuCauXuatHang, ChiTietYeuCauXuatHang, StockMovement)
    return ([model.objects.count() for model in models],
            list(TonKho.objects.order_by("pk").values_list("pk", "so_luong")))


class FormMarkup(HTMLParser):
    def __init__(self, html):
        super().__init__()
        self.buttons = []
        self.button_depth = 0
        self.nested_buttons = False
        self.links = []
        self.feed(html)

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == "button":
            self.nested_buttons |= self.button_depth > 0
            self.button_depth += 1
            self.buttons.append(attrs)
        elif tag == "a":
            self.links.append(attrs.get("href"))

    def handle_endtag(self, tag):
        if tag == "button":
            self.button_depth = max(0, self.button_depth - 1)


@pytest.mark.parametrize("kind,view", [
    ("import", "purchasing:import_create"),
    ("outbound", "sales_outbound:outbound_create"),
    ("stocktake", "stocktakes:stocktake_create"),
])
def test_real_download_can_be_uploaded_without_writing(client, users, sample_catalog, kind, view):
    client.force_login(users["mgr_cs1"])
    before = business_state()
    response = client.get(reverse("inventory:download_template", args=[kind]))
    assert response.status_code == 200
    # Use the exact workbook returned by the application, including its banner.
    upload = SimpleUploadedFile("downloaded.xlsx", response.content)
    result = client.post(reverse(view), {"form_action": "upload_excel", "excel_file": upload})
    assert result.status_code == 200
    preview = result.context["excel_preview"]
    assert preview["valid"], preview["errors"]
    assert {row["variant_id"] for row in preview["rows"]} == {
        sample_catalog["v_m"].id, sample_catalog["v_l"].id,
    }
    assert preview["rows"][0]["row_number"] == 5
    assert business_state() == before


@pytest.mark.parametrize("kind,headers,row", [
    ("import", ["Mã SKU", "Số lượng", "Đơn giá"], ["SP001-DEN-M", 2, 100000]),
    ("outbound", ["Mã SKU", "Số lượng"], ["SP001-DEN-M", 2]),
    ("stocktake", ["Mã SKU", "Số lượng thực tế"], ["SP001-DEN-M", 0]),
])
def test_headers_at_first_or_fourth_row_and_normalized_unicode(sample_catalog, branches, kind, headers, row):
    for header_row in (1, 4):
        # NFD accents, extra whitespace and NBSP come from ordinary Excel editing.
        edited = [unicodedata.normalize("NFD", "  " + h.replace(" ", "\u00a0  ").upper() + "  ") for h in headers]
        result = parse(kind, workbook_upload(edited, [row], header_row), branches["CS1"])
        assert result["valid"], result["errors"]
        assert result["rows"][0]["row_number"] == header_row + 1


@pytest.mark.parametrize("kind,headers,row", [
    ("import", ["Mã SKU", "Số lượng", "Đơn giá"], ["UNKNOWN", 2, 100000]),
    ("outbound", ["Mã SKU", "Số lượng"], ["SP001-DEN-M", 31]),
    ("stocktake", ["Mã SKU", "Số lượng thực tế"], ["SP001-DEN-M", -1]),
])
def test_errors_keep_physical_excel_row_number(sample_catalog, branches, kind, headers, row):
    result = parse(kind, workbook_upload(headers, [row], 4), branches["CS1"])
    assert not result["valid"]
    assert result["file_error"] is None
    assert result["rows"][0]["row_number"] == 5
    assert any("Dòng 5:" in error for error in result["errors"])
    if kind == "outbound":
        assert any("vượt quá tồn kho" in error for error in result["errors"])


@pytest.mark.parametrize("kind,headers,row", [
    ("import", ["Mã SKU", "Số lượng", "Đơn giá"], ["SP001-DEN-M", 2, 100000]),
    ("outbound", ["Mã SKU", "Số lượng"], ["SP001-DEN-M", 2]),
    ("stocktake", ["Mã SKU", "Số lượng thực tế"], ["SP001-DEN-M", 0]),
])
def test_missing_and_duplicate_required_headers_rejected(sample_catalog, branches, kind, headers, row):
    missing = parse(kind, workbook_upload(headers[1:], [row[1:]]), branches["CS1"])
    assert not missing["valid"] and missing["file_error"]
    duplicate = parse(kind, workbook_upload(headers + ["Mã hàng"], [row + [row[0]]]), branches["CS1"])
    assert not duplicate["valid"] and duplicate["file_error"]


def test_stocktake_legacy_quantity_alias_and_actual_quantity_priority(sample_catalog, branches):
    for headers, values in [
        (["Mã hàng", "Số lượng"], ["SP001-DEN-M", 0]),
        (["Mã SKU", "Số lượng", "Số lượng thực tế"], ["SP001-DEN-M", 99, 0]),
        (["Ghi chú", "Số lượng thực tế", "Mã SKU"], ["Đã đếm", 0, "SP001-DEN-M"]),
    ]:
        result = parse_stocktake_excel(workbook_upload(headers, [values], 4), branches["CS1"])
        assert result["valid"], result["errors"]
        row = result["rows"][0]
        assert row["so_luong_thuc_te"] == 0
        assert row["so_luong_he_thong"] == 30
        assert row["chenh_lech"] == -30


def test_request_creation_roles_ui_http_and_service(client, users, branches, sample_catalog):
    url = reverse("sales_outbound:request_list")
    for username in ("admin", "mgr_cs1", "mgr_cs2", "staff_cs1", "staff_cs2"):
        client.force_login(users[username])
        is_staff = username.startswith("staff")
        branch = branches["CS2"] if username.endswith("cs2") else branches["CS1"]
        page = client.get(reverse("sales_outbound:outbound_list"))
        assert (url in FormMarkup(page.content.decode()).links) == is_staff
        requests = client.get(url)
        assert requests.status_code == 200
        assert ('id="reqLinesContainer"' in requests.content.decode()) == is_staff
        before = (YeuCauXuatHang.objects.count(), ThongBao.objects.count())
        payload = {
            "ten_khach_hang": "Khách kiểm thử", "phuong_thuc_giao": "COUNTER",
            "line_variant_id": [sample_catalog["v_m"].id], "line_so_luong": [1],
        }
        result = client.post(url, payload)
        if is_staff:
            assert result.status_code == 302
            assert YeuCauXuatHang.objects.count() == before[0] + 1
            assert YeuCauXuatHang.objects.latest("pk").co_so == branch
        else:
            assert result.status_code == 403
            assert (YeuCauXuatHang.objects.count(), ThongBao.objects.count()) == before
            with pytest.raises(PermissionDenied):
                create_outbound_request(actor=users[username], co_so=branch,
                    line_specs=[OutboundRequestLineSpec(sample_catalog["v_m"].id, 1)])


def test_manager_keeps_processing_requests_and_cross_branch_guard(client, users, branches, sample_catalog):
    req = create_outbound_request(actor=users["staff_cs2"], co_so=branches["CS2"],
        line_specs=[OutboundRequestLineSpec(sample_catalog["v_m"].id, 1)])
    client.force_login(users["mgr_cs1"])
    response = client.get(reverse("sales_outbound:request_list"))
    assert req not in list(response.context["requests"])
    forbidden = client.post(reverse("sales_outbound:request_reject", args=[req.id]))
    assert forbidden.status_code == 403
    req.refresh_from_db()
    assert req.trang_thai == YeuCauXuatHang.Status.PENDING

    client.force_login(users["mgr_cs2"])
    response = client.get(reverse("sales_outbound:request_list"))
    assert req in list(response.context["requests"])
    assert 'id="reqLinesContainer"' not in response.content.decode()
    create_page = client.get(reverse("sales_outbound:outbound_create"), {"from_request": req.id})
    assert create_page.status_code == 200
    assert create_page.context["initial_lines"][0]["variant_id"] == sample_catalog["v_m"].id
    rejected = client.post(reverse("sales_outbound:request_reject", args=[req.id]), {"reason": "Khách hủy"})
    assert rejected.status_code == 302
    assert client.get(rejected.url).status_code == 200
    req.refresh_from_db()
    assert req.trang_thai == YeuCauXuatHang.Status.REJECTED
    assert ThongBao.objects.filter(nguoi_nhan=users["staff_cs2"], related_entity_id=req.id).exists()


def test_staff_cannot_create_request_for_another_branch(users, branches, sample_catalog):
    with pytest.raises(PermissionDenied):
        create_outbound_request(actor=users["staff_cs1"], co_so=branches["CS2"],
            line_specs=[OutboundRequestLineSpec(sample_catalog["v_m"].id, 1)])


def test_admin_request_processing_respects_all_and_selected_branch(client, users, branches, sample_catalog):
    requests = {}
    for branch_code, staff in (("CS1", "staff_cs1"), ("CS2", "staff_cs2")):
        requests[branch_code] = create_outbound_request(
            actor=users[staff], co_so=branches[branch_code],
            line_specs=[OutboundRequestLineSpec(sample_catalog["v_m"].id, 1)],
        )
    client.force_login(users["admin"])
    url = reverse("sales_outbound:request_list")
    all_branches = client.get(url, {"branch": "ALL"})
    assert all_branches.status_code == 200
    assert set(all_branches.context["requests"]) == set(requests.values())
    assert client.session["selected_branch_id"] == "ALL"
    assert set(client.get(url).context["requests"]) == set(requests.values())
    for branch_code in ("CS1", "CS2"):
        selected = client.get(url, {"branch": branch_code})
        assert selected.status_code == 200
        assert list(selected.context["requests"]) == [requests[branch_code]]
        assert client.session["selected_branch_id"] == branches[branch_code].id
        assert 'id="reqLinesContainer"' not in selected.content.decode()


def test_stocktake_has_one_save_button_and_separate_stock_column(client, users, sample_catalog):
    url = reverse("stocktakes:stocktake_create")
    for role in ("admin", "mgr_cs1", "mgr_cs2"):
        client.force_login(users[role])
        response = client.get(url)
        assert response.status_code == 200
        html = response.content.decode()
        markup = FormMarkup(html)
        assert not markup.nested_buttons
        assert len([button for button in markup.buttons if button.get("name") == "form_action" and button.get("value") == "submit"]) == 1
        assert "[Hệ thống:" not in html
        assert "Tồn hệ thống tại" in html
        assert "Số lượng thực tế" in html


def test_stocktake_validation_preserves_counted_quantity(client, users, sample_catalog):
    client.force_login(users["mgr_cs1"])
    response = client.post(reverse("stocktakes:stocktake_create"), {
        "form_action": "submit", "nguoi_kiem_ke": "", "muc_dich": "Đếm cuối ngày",
        "line_variant_id": [sample_catalog["v_m"].id], "line_so_luong_thuc_te": [7],
    })
    assert response.status_code == 200
    assert response.context["initial_lines"][0]["so_luong_thuc_te"] == "7"
    html = response.content.decode()
    assert 'value="7"' in html
    assert "[Hệ thống:" not in html
    assert not KiemKe.objects.exists()
